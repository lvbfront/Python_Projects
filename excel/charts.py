import openpyxl
from pathlib import Path


excelFile = openpyxl.load_workbook(Path.home() / Path("C:\\Users\\abdu4\\OneDrive\\سطح المكتب\\python course\\excel", "example.xlsx"))
print(excelFile.sheetnames)
sheet = excelFile['Sheet1']


title = openpyxl.chart.Reference(sheet, min_col = 1, max_col = 1, min_row = 1, max_row = 6)
data = openpyxl.chart.Reference(sheet, min_col = 2, max_col = 2, min_row = 1, max_row = 6)

chart = openpyxl.chart.LineChart()

chart.title = 'my chart'
chart.add_data(data= data)
chart.set_categories(title)

sheet.add_chart(chart, 'N8')

excelFile.save(filename = Path.home() / Path("C:\\Users\\abdu4\\OneDrive\\سطح المكتب\\python course\\excel", "example.xlsx"))